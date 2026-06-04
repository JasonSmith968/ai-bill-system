"""PDF and Excel report generator using reportlab and openpyxl."""

import os
import logging
from datetime import datetime

logger = logging.getLogger(__name__)


class ReportGenerator:
    """Generate PDF and Excel reports from BI data."""

    def __init__(self, user_id, bi_data, username='User'):
        self.user_id = user_id
        self.bi_data = bi_data
        self.username = username

    def generate_pdf(self, report_type, output_path):
        """Generate a styled PDF report."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        )
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        # Try to register Chinese font
        font_name = 'Helvetica'  # fallback
        font_search_paths = [
            'SimHei.ttf',                           # current dir
            '/usr/share/fonts/truetype/SimHei.ttf', # Linux
            'C:/Windows/Fonts/simhei.ttf',          # Windows
            '/System/Library/Fonts/SimHei.ttf',     # macOS
            os.path.join(os.path.dirname(__file__), '..', 'fonts', 'SimHei.ttf'),
        ]
        for fpath in font_search_paths:
            try:
                pdfmetrics.registerFont(TTFont('SimHei', fpath))
                font_name = 'SimHei'
                break
            except Exception:
                continue

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc = SimpleDocTemplate(output_path, pagesize=A4,
                                topMargin=20 * mm, bottomMargin=20 * mm)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('CNTitle', parent=styles['Title'],
                                     fontName=font_name, fontSize=20, spaceAfter=12)
        heading_style = ParagraphStyle('CNHeading', parent=styles['Heading2'],
                                       fontName=font_name, fontSize=14, spaceAfter=8)
        body_style = ParagraphStyle('CNBody', parent=styles['Normal'],
                                    fontName=font_name, fontSize=10, spaceAfter=6)

        story = []
        now = datetime.now().strftime('%Y-%m-%d %H:%M')
        report_label = {'weekly': '周报', 'monthly': '月报', 'custom': '自定义报告'}.get(report_type, '报告')

        # Title page
        story.append(Paragraph(f'AI 智能记账 — 财务{report_label}', title_style))
        story.append(Paragraph(f'用户: {self.username} | 生成时间: {now}', body_style))
        story.append(Spacer(1, 20 * mm))

        # Health score
        health = self.bi_data.get('health', {})
        if not isinstance(health, dict) or 'error' not in health:
            score = health.get('score', 0)
            level = health.get('level', 'N/A')
            story.append(Paragraph(f'财务健康评分: {score}/100 ({level})', heading_style))
            for rec in health.get('recommendations', []):
                story.append(Paragraph(f'• {rec}', body_style))
            story.append(Spacer(1, 10 * mm))

        # Forecast
        forecast = self.bi_data.get('forecast', {})
        if not isinstance(forecast, dict) or 'error' not in forecast:
            story.append(Paragraph('下月预测', heading_style))
            story.append(Paragraph(
                f"预测支出: ¥{forecast.get('predicted_expense', 0):,.2f} | "
                f"预测收入: ¥{forecast.get('predicted_income', 0):,.2f} | "
                f"趋势: {forecast.get('trend', 'N/A')} | "
                f"置信度: {forecast.get('confidence', 0) * 100:.0f}%",
                body_style
            ))
            story.append(Spacer(1, 10 * mm))

        # Category composition
        comp = self.bi_data.get('category_composition', {})
        cats = comp.get('categories', [])
        if cats:
            story.append(Paragraph('支出分类构成', heading_style))
            tdata = [['分类', '金额', '占比', '环比变化']]
            for c in cats[:10]:
                arrow = '↑' if c.get('mom_change', 0) > 0 else '↓'
                tdata.append([
                    c['name'],
                    f"¥{c['amount']:,.2f}",
                    f"{c['percentage']:.1f}%",
                    f"{arrow} {abs(c.get('mom_change_pct', 0)):.1f}%"
                ])
            t = Table(tdata, colWidths=[100, 100, 60, 80])
            t.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), font_name),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#334155')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
                ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ]))
            story.append(t)
            story.append(Spacer(1, 10 * mm))

        # Anomalies
        anomalies_data = self.bi_data.get('anomalies', {})
        anomalies = anomalies_data.get('anomalies', [])
        if anomalies:
            story.append(Paragraph('异常消费告警', heading_style))
            for a in anomalies[:5]:
                severity = '高' if a.get('severity') == 'high' else '中'
                story.append(Paragraph(
                    f"[{severity}] {a.get('date', '')} {a.get('category', '')} "
                    f"¥{a.get('amount', 0):,.2f} — {a.get('reason', '')}",
                    body_style
                ))
            story.append(Spacer(1, 10 * mm))

        # Cash flow
        cf = self.bi_data.get('cash_flow', {})
        if cf.get('labels'):
            story.append(Paragraph('现金流概览', heading_style))
            cf_data = [['月份', '收入', '支出', '净现金流']]
            for i, label in enumerate(cf['labels']):
                cf_data.append([
                    label,
                    f"¥{cf['income'][i]:,.2f}",
                    f"¥{cf['expense'][i]:,.2f}",
                    f"¥{cf['net'][i]:,.2f}"
                ])
            t = Table(cf_data, colWidths=[80, 100, 100, 100])
            t.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (-1, -1), font_name),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#334155')),
                ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ]))
            story.append(t)

        doc.build(story)
        return output_path

    def generate_enhanced_excel(self, output_path):
        """Generate Excel with multiple sheets and charts."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.chart import BarChart, PieChart, LineChart, Reference

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = Workbook()

        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1'),
        )

        # Sheet 1: Overview
        ws = wb.active
        ws.title = '财务概览'
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18

        health = self.bi_data.get('health', {})
        forecast = self.bi_data.get('forecast', {})

        overview_data = [
            ['指标', '数值'],
            ['财务健康评分', f"{health.get('score', 0)}/100"],
            ['预测下月支出', f"¥{forecast.get('predicted_expense', 0):,.2f}"],
            ['预测下月收入', f"¥{forecast.get('predicted_income', 0):,.2f}"],
            ['支出趋势', forecast.get('trend', 'N/A')],
        ]
        for row in overview_data:
            ws.append(row)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Sheet 2: Category Breakdown
        ws2 = wb.create_sheet('分类构成')
        ws2.column_dimensions['A'].width = 15
        ws2.column_dimensions['B'].width = 15
        ws2.column_dimensions['C'].width = 10
        ws2.column_dimensions['D'].width = 12

        comp = self.bi_data.get('category_composition', {})
        cats = comp.get('categories', [])
        ws2.append(['分类', '金额', '占比(%)', '环比变化(%)'])
        for c in cats:
            ws2.append([c['name'], c['amount'], c['percentage'], c.get('mom_change_pct', 0)])
        for cell in ws2[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        for row in ws2.iter_rows(min_row=2, max_row=len(cats) + 1):
            for cell in row:
                cell.border = thin_border

        # Pie chart for categories
        if cats:
            pie = PieChart()
            pie.title = '支出分类构成'
            labels = Reference(ws2, min_col=1, min_row=2, max_row=len(cats) + 1)
            data = Reference(ws2, min_col=2, min_row=1, max_row=len(cats) + 1)
            pie.add_data(data, titles_from_data=True)
            pie.set_categories(labels)
            pie.width = 16
            pie.height = 12
            ws2.add_chart(pie, 'F2')

        # Sheet 3: Cash Flow
        ws3 = wb.create_sheet('现金流')
        ws3.column_dimensions['A'].width = 12
        ws3.column_dimensions['B'].width = 15
        ws3.column_dimensions['C'].width = 15
        ws3.column_dimensions['D'].width = 15
        ws3.column_dimensions['E'].width = 15

        cf = self.bi_data.get('cash_flow', {})
        ws3.append(['月份', '收入', '支出', '净现金流', '累计'])
        for i, label in enumerate(cf.get('labels', [])):
            ws3.append([
                label,
                cf['income'][i] if i < len(cf['income']) else 0,
                cf['expense'][i] if i < len(cf['expense']) else 0,
                cf['net'][i] if i < len(cf['net']) else 0,
                cf['cumulative'][i] if i < len(cf['cumulative']) else 0,
            ])
        for cell in ws3[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Line chart for cash flow
        if cf.get('labels'):
            line = LineChart()
            line.title = '现金流趋势'
            line.width = 18
            line.height = 12
            labels_ref = Reference(ws3, min_col=1, min_row=2, max_row=len(cf['labels']) + 1)
            for col_idx in [2, 3, 4]:
                data_ref = Reference(ws3, min_col=col_idx, min_row=1, max_row=len(cf['labels']) + 1)
                line.add_data(data_ref, titles_from_data=True)
            line.set_categories(labels_ref)
            ws3.add_chart(line, 'G2')

        wb.save(output_path)
        return output_path
